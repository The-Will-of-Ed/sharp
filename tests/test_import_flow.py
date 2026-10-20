"""Real import-to-PDF plumbing; agent review is supplied, never simulated as AI."""

import asyncio
import json
import os
import subprocess
import sys

import pytest


def test_multiple_resume_folders_through_cli_and_real_mcp(tmp_path):
    pytest.importorskip("mcp")
    from docx import Document
    from mcp import ClientSession
    from mcp.client.stdio import StdioServerParameters, stdio_client
    from pypdf import PdfReader
    from reportlab.pdfgen.canvas import Canvas

    claim = "Built a Python API that reduced weekly manual reporting from three hours to one."
    resumes = tmp_path / "sources/resumes"
    notes = tmp_path / "sources/notes"
    resumes.mkdir(parents=True)
    notes.mkdir(parents=True)
    pdf = Canvas(str(resumes / "resume.pdf"))
    pdf.drawString(40, 750, claim)
    pdf.save()
    document = Document()
    document.add_paragraph(claim)
    document.add_paragraph("The same contribution appears in both sources.")
    document.save(notes / "resume-notes.docx")

    cli_output = tmp_path / "intake.json"
    subprocess.run(
        [
            sys.executable,
            "-m",
            "sharp",
            "intake",
            str(resumes),
            str(notes),
            "--output",
            str(cli_output),
        ],
        check=True,
        capture_output=True,
    )
    assert len(json.loads(cli_output.read_text(encoding="utf-8"))) == 2

    def unpack(result):
        assert not result.is_error, str(result)
        value = result.structured_content
        if value is None:
            values = [json.loads(c.text) for c in result.content if c.type == "text"]
            value = values[0] if len(values) == 1 else values
        return value.get("result", value) if isinstance(value, dict) else value

    async def run():
        server = StdioServerParameters(
            command=sys.executable,
            args=["-m", "sharp", "mcp", "--root", str(tmp_path)],
            env=os.environ.copy(),
        )
        async with stdio_client(server) as (read, write):
            async with ClientSession(read, write) as client:
                await client.initialize()
                unpack(await client.call_tool("workspace_initialize", {}))
                current = unpack(await client.call_tool("career_inspect", {}))
                extracted = unpack(
                    await client.call_tool(
                        "material_extract",
                        {
                            "paths": ["sources/resumes", "sources/notes"],
                        },
                    )
                )
                assert len(extracted) == 2
                assert all(claim in item["text"] for item in extracted)
                assert all(item["source"] and item["content_hash"] for item in extracted)
                # A reviewed agent/user proposal consolidates the duplicate claim.
                # The extraction layer is not credited with semantic deduplication.
                proposed = {
                    "basics": {"name": "Avery Example", "email": "avery@example.com"},
                    "work": [{"x_id": "work1", "name": "Example Systems", "position": "Engineer"}],
                    "x_resume": {
                        "evidence": [
                            {"id": "source1", "source": extracted[0]["source"]},
                            {"id": "source2", "source": extracted[1]["source"]},
                        ],
                        "bullets": [
                            {
                                "id": "api",
                                "owner_id": "work1",
                                "text": claim,
                                "evidence_ids": ["source1", "source2"],
                                "approval": "approved",
                            }
                        ],
                    },
                }
                unpack(
                    await client.call_tool(
                        "career_replace",
                        {
                            "career_path": "master.yaml",
                            "proposed_career": proposed,
                            "expected_hash": current["hash"],
                        },
                    )
                )
                saved = unpack(await client.call_tool("career_inspect", {}))["career"]
                assert len(saved["x_resume"]["bullets"]) == 1
                assert saved["x_resume"]["bullets"][0]["confidence"]["level"] == "unrated"
                result = unpack(
                    await client.call_tool(
                        "document_create",
                        {
                            "request": {
                                "career_path": "master.yaml",
                                "settings_path": "settings.local.yaml",
                                "output_dir": "outputs/import-demo",
                                "title": "Software Engineer",
                            },
                            "choice_plan": {"ranked": [{"candidate_id": "api"}]},
                        },
                    )
                )
                assert result["pages"] == 1
                assert result["selected_ids"] == ["api"]
                pdf = next((tmp_path / "outputs/import-demo").glob("*.pdf"))
                assert claim in PdfReader(pdf).pages[0].extract_text()
                # Complete replacement still refuses a stale revision.
                stale = await client.call_tool(
                    "career_replace",
                    {
                        "career_path": "master.yaml",
                        "proposed_career": proposed,
                        "expected_hash": current["hash"],
                    },
                )
                assert stale.is_error

    asyncio.run(run())
