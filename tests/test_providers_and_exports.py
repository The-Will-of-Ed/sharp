import json
import pytest
from sharp.adapters.providers import LiteLLMChooser, AgentChooser, choice_prompt
from sharp.career import choice_input, prepare_document
from sharp.layout import fit
from sharp.adapters.pdf import PdfRenderer


def test_one_generative_call_with_confidence_and_no_fake_reasons(career, settings, plan):
    calls = []
    def transport(route, data):
        calls.append((route, data))
        return {"choices": [{"message": {"content": plan.model_dump_json()}}]}
    chooser = LiteLLMChooser("http://localhost:4000", "test", transport=transport)
    result = chooser.choose(choice_input(career, settings, "Python developer"))
    assert len(calls) == 1
    assert "very_high" in calls[0][1]["messages"][1]["content"]
    assert result.usage == {}
    assert result.ranked[0].reason is None


def test_decision_models_rank_without_rewriting(career, settings):
    def transport(route, data):
        assert route == "/v1/decisions"
        assert all(q["type"] == "score" for q in data["questions"])
        return {"answers": [{"name": q["name"], "type": "score", "score": 1} for q in data["questions"]]}
    result = LiteLLMChooser("http://localhost:4000", "jev", mode="decisions", transport=transport).choose(choice_input(career, settings, "Python"))
    assert result.ranked
    assert all(d.reason is None for d in result.ranked)
    assert result.summary_id == "s1"


def test_decision_refusal_is_not_low_score(career, settings):
    def transport(route, data): return {"answers": [{"name": data["questions"][0]["name"], "type": "refusal"}]}
    with pytest.raises(ValueError, match="refus|missing"):
        LiteLLMChooser("http://localhost:4000", "jev", mode="decisions", transport=transport).choose(choice_input(career, settings, "Python"))


def test_agent_uses_same_contract(career, settings, plan):
    calls = []
    def ask(prompt):
        calls.append(prompt)
        return plan.model_dump_json()
    assert AgentChooser(ask).choose(choice_input(career, settings, "Python")).ranked == plan.ranked
    assert len(calls) == 1


@pytest.mark.parametrize("format", ["pptx", "docx"])
def test_editable_exports_keep_text_and_default_regular(tmp_path, career, settings, plan, format):
    from sharp.adapters.office import OfficeExporter
    layout = fit(prepare_document(career, settings, plan), settings, plan, PdfRenderer())
    path = tmp_path / f"resume.{format}"
    artifact = OfficeExporter(format).export(layout, settings, str(path))
    assert artifact.checks["editable_text"]
    if format == "pptx":
        from pptx import Presentation
        prs = Presentation(path)
        text = "\n".join(shape.text for slide in prs.slides for shape in slide.shapes if shape.has_text_frame)
        assert all(not run.font.bold for slide in prs.slides for shape in slide.shapes if shape.has_text_frame for p in shape.text_frame.paragraphs for run in p.runs)
    else:
        from docx import Document
        doc = Document(path)
        text = "\n".join(p.text for p in doc.paragraphs)
        assert all(not r.bold for p in doc.paragraphs for r in p.runs)
    assert "Alex Example" in text and "not realized" in text
