"""Local files and ingestion; never infers facts, approval, or confidence."""

from __future__ import annotations

import hashlib
import json
import os
from contextlib import contextmanager
from pathlib import Path

import yaml

from sharp.career import digest, validate_career
from sharp.contracts import Career, ChoicePlan, Customization, Settings


def load(path, model):
    return model.model_validate(yaml.safe_load(Path(path).read_text(encoding="utf-8")))


def save(path, value):
    target = Path(path)
    target.parent.mkdir(parents=True, exist_ok=True)
    data = value.model_dump(mode="json", by_alias=True) if hasattr(value, "model_dump") else value
    temporary = target.with_suffix(target.suffix + ".tmp")
    temporary.write_text(json.dumps(data, indent=2, ensure_ascii=False), encoding="utf-8")
    temporary.replace(target)


@contextmanager
def lock(path):
    target = Path(str(path) + ".lock")
    try:
        fd = os.open(target, os.O_CREAT | os.O_EXCL | os.O_WRONLY)
    except FileExistsError as exc:
        raise ValueError("File is being updated; retry after the current update") from exc
    try:
        yield
    finally:
        os.close(fd)
        target.unlink(missing_ok=True)


class SettingsFile:
    def __init__(self, path):
        self.path = Path(path)

    def initialize(self, settings=None):
        self.path.parent.mkdir(parents=True, exist_ok=True)
        with lock(self.path):
            if self.path.exists():
                raise ValueError(
                    "Local settings already exist; initialization never overwrites them"
                )
            save(self.path, settings or Settings())
        return self.read()

    def read(self):
        return load(self.path, Settings)

    def update(self, changes, *, expected_hash, apply_everywhere=False):
        with lock(self.path):
            current = self.read()
            if digest(current) != expected_hash:
                raise ValueError("Settings revision changed; inspect and retry")
            values = current.model_dump()
            for key, value in changes.items():
                parts, target = key.split("."), values
                for part in parts[:-1]:
                    if part not in target:
                        target[part] = {}
                    target = target[part]
                    if not isinstance(target, dict):
                        raise ValueError(f"Not a settings object: {key}")
                target[parts[-1]] = value
                if apply_everywhere and key.startswith("styles."):
                    for exception in values["exceptions"].values():
                        exception.pop(parts[-1], None)
            result = Settings.model_validate(values)
            save(self.path, result)
            return result

    def customize(
        self,
        action,
        *,
        expected_hash,
        customization=None,
        id=None,
        runner=None,
        preview_payload=None,
    ):
        current = self.read()
        if digest(current) != expected_hash:
            raise ValueError("Settings revision changed")
        functions = list(current.customizations)
        existing = {f.id: f for f in functions}
        if action in ("install", "update"):
            rule = Customization.model_validate(customization)
            if action == "install" and rule.id in existing:
                raise ValueError("Customization already exists")
            if action == "update" and rule.id not in existing:
                raise ValueError("Customization does not exist")
            if rule.id in existing and rule.version <= existing[rule.id].version:
                raise ValueError("Update needs a higher customization version")
            if runner is None:
                raise ValueError("Install/update requires restricted runtime compilation")
            runner.install(rule)
            if rule.enabled:
                if preview_payload is None:
                    raise ValueError(
                        "Enabled rules need representative preview_payload before activation"
                    )
                from sharp.contracts import Changes, Decorations

                output = runner.run(rule, preview_payload)
                model = {
                    "selection": ChoicePlan,
                    "before_layout": Changes,
                    "after_layout": Decorations,
                }[rule.stage]
                model.model_validate(output)
            functions = (
                [rule if f.id == rule.id else f for f in functions]
                if rule.id in existing
                else [*functions, rule]
            )
        elif action == "remove":
            if id not in existing:
                raise ValueError("Customization does not exist")
            functions = [f for f in functions if f.id != id]
        else:
            raise ValueError("Expected install, update, or remove")
        return self.update(
            {"customizations": [f.model_dump() for f in functions]},
            expected_hash=expected_hash,
        )

    def portable(self):
        settings = self.read()
        fonts = [s.font for s in settings.styles.values()] + [
            e.get("font", "") for e in settings.exceptions.values()
        ]
        if any("/" in font or "\\" in font or ":" in font for font in fonts):
            raise ValueError(
                "Portable settings cannot contain local font paths; choose a portable font first"
            )
        return settings.model_dump()


def import_material(paths: list[str], *, allowed_root=None) -> list[dict]:
    files = []
    for entry in paths:
        path = Path(entry)
        files.extend(
            sorted(
                p
                for p in path.rglob("*")
                if p.is_file()
                and not any(part.startswith(".") for part in p.relative_to(path).parts)
            )
            if path.is_dir()
            else [path]
        )
    result = []
    for path in files:
        if allowed_root is not None and not path.resolve().is_relative_to(
            Path(allowed_root).resolve()
        ):
            raise ValueError("Source file resolves outside the configured workspace")
        suffix = path.suffix.lower()
        if suffix not in (
            ".txt",
            ".md",
            ".yaml",
            ".yml",
            ".json",
            ".pdf",
            ".docx",
            ".pptx",
        ):
            continue
        if suffix == ".pdf":
            from pypdf import PdfReader

            text = "\n\n".join(
                f"Page {i + 1}\n{p.extract_text() or ''}"
                for i, p in enumerate(PdfReader(path).pages)
            )
        elif suffix == ".docx":
            from docx import Document

            document = Document(path)
            text = "\n".join(p.text for p in document.paragraphs)
            text += "\n" + "\n".join(
                " | ".join(c.text for c in row.cells) for t in document.tables for row in t.rows
            )
        elif suffix == ".pptx":
            from pptx import Presentation

            text = "\n\n".join(
                f"Slide {i + 1}\n" + "\n".join(s.text for s in slide.shapes if s.has_text_frame)
                for i, slide in enumerate(Presentation(path).slides)
            )
        else:
            text = path.read_text(encoding="utf-8-sig")
        result.append(
            {
                "source": str(path.resolve()),
                "content_hash": hashlib.sha256(path.read_bytes()).hexdigest(),
                "text": text,
                "review": "Extracted text only; review original visuals and user context before proposing claims.",
            }
        )
    return result


def export_json_resume(career: Career, plan: ChoicePlan | None = None):
    validate_career(career)
    data = career.model_dump(mode="json", by_alias=True, exclude_none=True)
    data.pop("x_resume", None)
    data.pop("$schema", None)
    decisions = {d.candidate_id: d for d in plan.ranked} if plan else None
    for kind in ("work", "projects", "volunteer"):
        for record in data[kind]:
            id = record.pop("x_id")
            record["highlights"] = []
            for b in career.x_resume.bullets:
                if (
                    b.owner_id != id
                    or b.approval != "approved"
                    or decisions is not None
                    and b.id not in decisions
                ):
                    continue
                variant = decisions[b.id].variant_id if decisions is not None else None
                record["highlights"].append(
                    next((v.text for v in b.variants if v.id == variant), b.text)
                )
    return data


def validate_standard(career: Career):
    """Validate the standard core with the pinned upstream schema, offline."""
    from importlib.resources import files

    from jsonschema import Draft7Validator, FormatChecker

    schema = json.loads(
        files("sharp").joinpath("defaults/jsonresume.schema.json").read_text(encoding="utf-8")
    )
    errors = sorted(
        Draft7Validator(schema, format_checker=FormatChecker()).iter_errors(
            export_json_resume(career)
        ),
        key=lambda e: str(list(e.path)),
    )
    if errors:
        error = errors[0]
        raise ValueError(f"JSON Resume core {'.'.join(map(str, error.path))}: {error.message}")
    return career


def migrate_legacy(data: dict) -> Career:
    """Explicit, non-destructive importer for this repository's old career shape."""
    if "basics" in data:
        return validate_career(Career.model_validate(data))
    person = data.get("person", {})
    basics = {
        "name": person.get("name", ""),
        "email": person.get("email", ""),
        "phone": person.get("phone", ""),
        "profiles": [],
    }
    for key in ("linkedin", "github"):
        if person.get(key):
            url = person[key]
            basics["profiles"].append(
                {
                    "network": key,
                    "url": url if url.startswith("http") else "https://" + url,
                }
            )
    result = {
        "basics": basics,
        "work": [],
        "projects": [],
        "education": [],
        "x_resume": {"bullets": [], "evidence": [], "skills": [], "summaries": []},
    }
    for kind, key in (("work", "experience"), ("projects", "projects")):
        for i, old in enumerate(data.get(key, [])):
            owner = str(old.get("id") or f"{kind}_{i}")
            record = {
                "x_id": owner,
                "name": old.get("company") or old.get("name") or old.get("title") or owner,
            }
            if kind == "work":
                record["position"] = old.get("role") or old.get("title") or ""
            for new, keys in (
                ("startDate", ("start", "start_date")),
                ("endDate", ("end", "end_date")),
            ):
                value = next((old[k] for k in keys if old.get(k)), None)
                if value and value != "Present":
                    record[new] = str(value)
            result[kind].append(record)
            for j, bullet in enumerate(old.get("bullets", [])):
                text = bullet.get("text", "") if isinstance(bullet, dict) else str(bullet)
                id = (
                    bullet.get("id", f"{owner}_{j}") if isinstance(bullet, dict) else f"{owner}_{j}"
                )
                evidence_id = f"source:{id}"
                result["x_resume"]["evidence"].append(
                    {
                        "id": evidence_id,
                        "kind": "user_attested",
                        "source": "Explicit legacy master import",
                    }
                )
                result["x_resume"]["bullets"].append(
                    {
                        "id": id,
                        "owner_id": owner,
                        "text": text,
                        "evidence_ids": [evidence_id],
                        "approval": "proposed",
                    }
                )
    for i, old in enumerate(data.get("education", [])):
        if isinstance(old, dict):
            result["education"].append(
                {
                    "institution": old.get("institution", old.get("school", "")),
                    "area": old.get("field", ""),
                    "studyType": old.get("degree", ""),
                }
            )
    return validate_career(Career.model_validate(result))
