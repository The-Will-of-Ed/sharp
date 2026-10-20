"""Optional editable exports. Office text engines can reflow differently from PDF."""

from pathlib import Path

from sharp.contracts import Artifact

from .pdf import segments


class OfficeExporter:
    def __init__(self, format):
        if format not in ("pptx", "docx"):
            raise ValueError("Expected pptx or docx")
        self.format = format

    def export(self, layout, settings, destination):
        if self.format == "pptx":
            self._pptx(layout, settings, destination)
        else:
            self._docx(layout, settings, destination)
        return Artifact(
            path=str(Path(destination).resolve()),
            format=self.format,
            checks={
                "editable_text": True,
                "pages_in_layout": layout.pages,
                "layout_fidelity": "explicit line positions"
                if self.format == "pptx"
                else "flow layout; recheck in Word before submitting",
            },
        )

    def _pptx(self, layout, settings, destination):
        from pptx import Presentation
        from pptx.dml.color import RGBColor
        from pptx.util import Pt

        prs = Presentation()
        prs.slide_width, prs.slide_height = (
            Pt(settings.page.width),
            Pt(settings.page.height),
        )
        for page in range(1, layout.pages + 1):
            slide = prs.slides.add_slide(prs.slide_layouts[6])
            for p in (p for p in layout.placed if p.page == page):
                leading = p.measurement.height / len(p.measurement.lines)
                for i, line in enumerate(p.measurement.lines):
                    shape = slide.shapes.add_textbox(
                        Pt(p.x), Pt(p.y + i * leading), Pt(p.width + 1), Pt(leading + 2)
                    )
                    frame = shape.text_frame
                    frame.word_wrap = False
                    frame.margin_top = frame.margin_bottom = frame.margin_left = (
                        frame.margin_right
                    ) = 0
                    paragraph = frame.paragraphs[0]
                    paragraph.space_before = paragraph.space_after = Pt(0)
                    for text, style in segments(p.block, line.start, line.end):
                        run = paragraph.add_run()
                        run.text = text
                        run.font.name = {
                            "Helvetica": "Arial",
                            "Times-Roman": "Times New Roman",
                            "Courier": "Courier New",
                        }.get(style.font, style.font)
                        run.font.size, run.font.bold, run.font.italic = (
                            Pt(style.size),
                            style.bold,
                            style.italic,
                        )
                        run.font.color.rgb = RGBColor.from_string(style.color[1:])
                        if p.block.url:
                            run.hyperlink.address = p.block.url
        prs.save(destination)

    def _docx(self, layout, settings, destination):
        from docx import Document
        from docx.shared import Pt, RGBColor

        doc = Document()
        section = doc.sections[0]
        section.page_width, section.page_height = (
            Pt(settings.page.width),
            Pt(settings.page.height),
        )
        top, right, bottom, left = settings.page.margins
        (
            section.top_margin,
            section.right_margin,
            section.bottom_margin,
            section.left_margin,
        ) = map(Pt, (top, right, bottom, left))
        normal = doc.styles["Normal"]
        normal.paragraph_format.space_before = normal.paragraph_format.space_after = Pt(0)
        current = 1
        for p in layout.placed:
            if p.page != current:
                doc.add_page_break()
                current = p.page
            paragraph = doc.add_paragraph()
            paragraph.paragraph_format.space_before = Pt(p.block.style.before)
            paragraph.paragraph_format.space_after = Pt(p.block.style.after)
            paragraph.paragraph_format.line_spacing = Pt(
                p.measurement.height / len(p.measurement.lines)
            )
            paragraph.paragraph_format.left_indent = Pt(p.block.style.indent)
            paragraph.paragraph_format.keep_with_next = p.block.kind in (
                "heading",
                "owner",
            )
            for text, style in segments(p.block, 0, len(p.block.text)):
                run = paragraph.add_run(text)
                run.bold, run.italic = style.bold, style.italic
                run.font.name = {
                    "Helvetica": "Arial",
                    "Times-Roman": "Times New Roman",
                    "Courier": "Courier New",
                }.get(style.font, style.font)
                run.font.size, run.font.color.rgb = (
                    Pt(style.size),
                    RGBColor.from_string(style.color[1:]),
                )
                if p.block.url:
                    from docx.opc.constants import RELATIONSHIP_TYPE
                    from docx.oxml import OxmlElement
                    from docx.oxml.ns import qn

                    relation = paragraph.part.relate_to(
                        p.block.url, RELATIONSHIP_TYPE.HYPERLINK, is_external=True
                    )
                    hyperlink = OxmlElement("w:hyperlink")
                    hyperlink.set(qn("r:id"), relation)
                    run._r.addprevious(hyperlink)
                    hyperlink.append(run._r)
        doc.save(destination)
